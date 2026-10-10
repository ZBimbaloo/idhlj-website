#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Convert Markdown files to styled Word (DOCX) documents."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------- color palette ----------
BRAND_ORANGE = RGBColor(0xF5, 0x7C, 0x00)   # Huanlejia brand
TEXT_DARK    = RGBColor(0x1F, 0x2A, 0x37)
TEXT_GREY    = RGBColor(0x55, 0x65, 0x6B)
TABLE_HEAD   = RGBColor(0xF5, 0x7C, 0x00)
TABLE_ALT    = RGBColor(0xFD, 0xF1, 0xE1)


# ---------- helpers ----------
def set_cell_bg(cell, rgb_hex: str):
    """Set cell background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), rgb_hex)
    tcPr.append(shd)


def set_run_font(run, name='Calibri', size=10.5, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)


def add_inline(paragraph, text: str, base_bold=False, base_size=10.5):
    """Parse inline markdown (bold/italic/code/link) and add runs."""
    # Split on inline patterns: **bold**, *italic*, `code`, [text](url)
    pattern = re.compile(
        r'(\*\*[^*]+\*\*)'           # **bold**
        r'|(\*[^*]+\*)'              # *italic*
        r'|(`[^`]+`)'                # `code`
        r'|(\[[^\]]+\]\([^)]+\))'    # [text](url)
    )
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r, size=base_size, bold=base_bold, color=TEXT_DARK)
        token = m.group(0)
        if token.startswith('**'):
            r = paragraph.add_run(token[2:-2])
            set_run_font(r, size=base_size, bold=True, color=TEXT_DARK)
        elif token.startswith('*'):
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, size=base_size, bold=base_bold, color=TEXT_DARK)
            r.italic = True
        elif token.startswith('`'):
            r = paragraph.add_run(token[1:-1])
            set_run_font(r, name='Consolas', size=base_size - 0.5,
                         bold=base_bold, color=RGBColor(0xC0, 0x39, 0x2B))
            rPr = r._element.get_or_add_rPr()
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), 'F2F2F2')
            rPr.append(shd)
        elif token.startswith('['):
            mm = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if mm:
                r = paragraph.add_run(mm.group(1))
                set_run_font(r, size=base_size, bold=base_bold,
                             color=RGBColor(0x05, 0x63, 0xC1))
                r.underline = True
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, size=base_size, bold=base_bold, color=TEXT_DARK)


def add_horizontal_rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'F5A623')
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_blockquote(doc, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.4)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '8')
    left.set(qn('w:color'), 'F5A623')
    pBdr.append(left)
    pPr.append(pBdr)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=TEXT_GREY)
    r.italic = True


# ---------- markdown parser ----------
def parse_table(lines, idx):
    """Parse a contiguous markdown table block. Return (rows, next_idx)."""
    rows = []
    while idx < len(lines) and lines[idx].lstrip().startswith('|'):
        raw = lines[idx].strip()
        cells = [c.strip() for c in raw.strip('|').split('|')]
        # Skip separator row (---)
        if all(re.match(r'^:?-+:?$', c) for c in cells if c):
            idx += 1
            continue
        rows.append(cells)
        idx += 1
    return rows, idx


def build_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = True

    for r_idx, row in enumerate(rows):
        is_header = (r_idx == 0)
        for c_idx in range(cols):
            cell = table.rows[r_idx].cells[c_idx]
            text = row[c_idx] if c_idx < len(row) else ''
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_header else WD_ALIGN_PARAGRAPH.LEFT
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            add_inline(p, text, base_bold=is_header,
                       base_size=9.5 if not is_header else 10)
            if is_header:
                set_cell_bg(cell, 'F57C00')
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.bold = True
            else:
                if r_idx % 2 == 0:
                    set_cell_bg(cell, 'FDF1E1')
    doc.add_paragraph()  # spacing after table


def add_heading(doc, text: str, level: int):
    sizes = {1: 22, 2: 16, 3: 13, 4: 11.5, 5: 10.5, 6: 10}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=sizes.get(level, 10), bold=True,
                 color=BRAND_ORANGE if level <= 2 else TEXT_DARK)
    if level == 1:
        p.paragraph_format.space_after = Pt(2)
        # bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '4')
        bottom.set(qn('w:color'), 'F5A623')
        pBdr.append(bottom)
        pPr.append(pBdr)


def render_markdown(doc, md_text: str):
    lines = md_text.splitlines()
    i = 0
    n = len(lines)
    in_code = False
    code_buf = []
    in_list = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith('```'):
            if not in_code:
                in_code = True
                code_buf = []
            else:
                # render code block
                p = doc.add_paragraph()
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'F4F4F4')
                pPr.append(shd)
                for c_line in code_buf:
                    r = p.add_run(c_line + '\n')
                    set_run_font(r, name='Consolas', size=9,
                                 color=RGBColor(0x33, 0x33, 0x33))
                in_code = False
                code_buf = []
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # blank line
        if not stripped:
            i += 1
            continue

        # horizontal rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            add_horizontal_rule(doc)
            i += 1
            continue

        # heading
        m = re.match(r'^(#{1,6})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            add_heading(doc, m.group(2).strip(), level)
            i += 1
            continue

        # table
        if stripped.startswith('|'):
            rows, i = parse_table(lines, i)
            build_table(doc, rows)
            continue

        # blockquote
        if stripped.startswith('>'):
            txt = stripped.lstrip('>').strip()
            add_blockquote(doc, txt)
            i += 1
            continue

        # ordered list
        m = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, m.group(2), base_size=10.5)
            i += 1
            continue

        # unordered list
        m = re.match(r'^[-*+]\s+(.*)$', stripped)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, m.group(1), base_size=10.5)
            i += 1
            continue

        # paragraph (collect contiguous non-empty non-special lines)
        para_lines = [stripped]
        j = i + 1
        while j < n:
            nxt = lines[j].strip()
            if (not nxt or nxt.startswith('#') or nxt.startswith('|')
                    or nxt.startswith('>') or nxt.startswith('-')
                    or nxt.startswith('*') or re.match(r'^\d+\.', nxt)
                    or nxt.startswith('```')
                    or re.match(r'^-{3,}$', nxt)):
                break
            para_lines.append(nxt)
            j += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.4
        add_inline(p, ' '.join(para_lines), base_size=10.5)
        i = j


# ---------- main ----------
def convert_file(md_path: Path, docx_path: Path, title: str):
    doc = Document()
    # default style
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = TEXT_DARK

    # page margins
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    # cover banner
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = cover.add_run('欢乐家印尼子公司官网 · 文档')
    set_run_font(r, size=10, bold=True, color=BRAND_ORANGE)
    doc.add_paragraph()

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tr = title_p.add_run(title)
    set_run_font(tr, size=26, bold=True, color=BRAND_ORANGE)
    title_p.paragraph_format.space_after = Pt(6)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    sr = sub_p.add_run(f'生成自：{md_path.name}')
    set_run_font(sr, size=10, color=TEXT_GREY)
    sr.italic = True

    add_horizontal_rule(doc)
    doc.add_paragraph()

    md_text = md_path.read_text(encoding='utf-8')
    render_markdown(doc, md_text)

    doc.save(docx_path)
    return docx_path


if __name__ == '__main__':
    base = Path(r'd:\code\idhlj-website')
    targets = [
        (base / '印尼官网需求文档.md',
         base / '印尼官网需求文档.docx',
         '欢乐家印尼子公司官网 · 需求文档'),
        (base / '素材缺失文档.md',
         base / '素材缺失文档.docx',
         '欢乐家印尼子公司官网 · 缺失素材汇总'),
    ]
    for src, dst, title in targets:
        if not src.exists():
            print(f'[skip] {src} not found')
            continue
        out = convert_file(src, dst, title)
        print(f'[ok]   {out}  ({out.stat().st_size:,} bytes)')
